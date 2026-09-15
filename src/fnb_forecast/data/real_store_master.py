"""Master data definitions and Excel parser for the real-world Vietnam cafe dataset."""

from pathlib import Path
from typing import Tuple
import openpyxl
import pandas as pd


REAL_STORE_ITEMS = [
    # Coffee (7 items)
    {"item_id": "CF_DEN", "item_name": "Cà phê đen", "category": "Coffee", "unit_price": 15000},
    {"item_id": "CF_SUA", "item_name": "Cà phê sữa", "category": "Coffee", "unit_price": 17000},
    {"item_id": "CF_BACXIU", "item_name": "Bạc xỉu", "category": "Coffee", "unit_price": 20000},
    {"item_id": "CF_TRUNG", "item_name": "Cà phê trứng", "category": "Coffee", "unit_price": 25000},
    {"item_id": "CF_MUOI", "item_name": "Cà phê muối", "category": "Coffee", "unit_price": 20000},
    {"item_id": "CF_BUONME", "item_name": "Cà phê Buôn Mê", "category": "Coffee", "unit_price": 25000},
    {"item_id": "CF_TIRAMISU", "item_name": "Cà phê Tiramisu", "category": "Coffee", "unit_price": 25000},

    # Topping (5 items)
    {"item_id": "TOP_TCDEN", "item_name": "Trân châu đen", "category": "Topping", "unit_price": 5000},
    {"item_id": "TOP_TCTRANG", "item_name": "Trân châu trắng", "category": "Topping", "unit_price": 5000},
    {"item_id": "TOP_KEMUOI", "item_name": "Kem muối", "category": "Topping", "unit_price": 7000},
    {"item_id": "TOP_KEMTRUNG", "item_name": "Kem trứng", "category": "Topping", "unit_price": 10000},
    {"item_id": "TOP_KEMBUONME", "item_name": "Kem Buôn Mê", "category": "Topping", "unit_price": 10000},

    # Olong / Soda / Tea (11 items)
    {"item_id": "TEA_DAU", "item_name": "Trà Olong Dâu", "category": "Tea_Soda", "unit_price": 20000},
    {"item_id": "TEA_VIETQUAT", "item_name": "Trà Olong Việt Quất", "category": "Tea_Soda", "unit_price": 20000},
    {"item_id": "TEA_KIWI", "item_name": "Trà Olong Kiwi", "category": "Tea_Soda", "unit_price": 20000},
    {"item_id": "TEA_CHANHDAY", "item_name": "Trà Olong Chanh Dây", "category": "Tea_Soda", "unit_price": 20000},
    {"item_id": "TEA_LUUDO", "item_name": "Trà Olong Lựu Đỏ", "category": "Tea_Soda", "unit_price": 20000},
    {"item_id": "TEA_OIHONG", "item_name": "Trà Olong Ổi Hồng", "category": "Tea_Soda", "unit_price": 20000},
    {"item_id": "TEA_TAOXANH", "item_name": "Trà Olong Táo Xanh", "category": "Tea_Soda", "unit_price": 20000},
    {"item_id": "TEA_DAOCAMSA", "item_name": "Trà Đào Cam Sả", "category": "Tea_Soda", "unit_price": 20000},
    {"item_id": "TEA_DUALUOI", "item_name": "Trà Olong Dưa Lưới", "category": "Tea_Soda", "unit_price": 20000},
    {"item_id": "TEA_MANGCAU", "item_name": "Trà Mãng Cầu", "category": "Tea_Soda", "unit_price": 20000},
    {"item_id": "TEA_HONGTRA", "item_name": "Hồng Trà Truyền Thống", "category": "Tea_Soda", "unit_price": 15000},

    # Iceblend (5 items)
    {"item_id": "ICE_CF_CARAMEL", "item_name": "Coffee Caramel Đá Xay", "category": "Iceblend", "unit_price": 30000},
    {"item_id": "ICE_TRAICAY", "item_name": "Đá Xay Trái Cây", "category": "Iceblend", "unit_price": 30000},
    {"item_id": "ICE_MATCHA", "item_name": "Matcha Đá Xay", "category": "Iceblend", "unit_price": 30000},
    {"item_id": "ICE_OREO", "item_name": "Oreo Đá Xay", "category": "Iceblend", "unit_price": 30000},
    {"item_id": "ICE_MILO", "item_name": "Milo Dầm Trân Châu", "category": "Iceblend", "unit_price": 25000},

    # Milktea (9 items)
    {"item_id": "MT_TRUYENTHONG", "item_name": "Trà Sữa Truyền Thống", "category": "Milktea", "unit_price": 15000},
    {"item_id": "MT_DUONGDEN", "item_name": "Sữa Tươi Trân Châu Đường Đen", "category": "Milktea", "unit_price": 20000},
    {"item_id": "MT_THAI", "item_name": "Trà Sữa Thái Xanh / Đỏ", "category": "Milktea", "unit_price": 20000},
    {"item_id": "MT_DAU_VQ", "item_name": "Trà Sữa Dâu / Việt Quất", "category": "Milktea", "unit_price": 20000},
    {"item_id": "MT_MATCHA", "item_name": "Trà Sữa Matcha", "category": "Milktea", "unit_price": 20000},
    {"item_id": "MT_KHOAIMON", "item_name": "Trà Sữa Khoai Môn", "category": "Milktea", "unit_price": 20000},
    {"item_id": "MT_TAOXANH", "item_name": "Trà Sữa Táo Xanh", "category": "Milktea", "unit_price": 20000},
    {"item_id": "MT_COMDEO", "item_name": "Trà Sữa Cốm Dẻo", "category": "Milktea", "unit_price": 25000},
    {"item_id": "MT_OLONG_LAI", "item_name": "Olong Lài Sữa", "category": "Milktea", "unit_price": 20000},

    # Matcha / Cacao / Dessert (6 items)
    {"item_id": "MAT_LATTE", "item_name": "Matcha Latte", "category": "Matcha_Cacao", "unit_price": 25000},
    {"item_id": "MAT_GAU", "item_name": "Matcha Sữa Gấu", "category": "Matcha_Cacao", "unit_price": 30000},
    {"item_id": "MAT_FRUIT", "item_name": "Matcha Trái Cây", "category": "Matcha_Cacao", "unit_price": 30000},
    {"item_id": "CAC_LATTE", "item_name": "Cacao Latte", "category": "Matcha_Cacao", "unit_price": 25000},
    {"item_id": "CAC_GAU", "item_name": "Cacao Sữa Gấu", "category": "Matcha_Cacao", "unit_price": 30000},
    {"item_id": "DES_PANNACOTTA", "item_name": "Panna Cotta", "category": "Dessert", "unit_price": 15000},
]


REAL_STORE_INGREDIENTS = [
    {"ingredient_id": "ING_CAFE", "name": "Cà phê pha phin / espresso", "base_unit": "ml", "pack_size": 1000.0, "unit_cost": 72.0, "lead_time_days": 2, "review_period_days": 7, "shelf_life_days": 14},
    {"ingredient_id": "ING_SUADAC", "name": "Sữa đặc", "base_unit": "g", "pack_size": 1000.0, "unit_cost": 65.0, "lead_time_days": 2, "review_period_days": 7, "shelf_life_days": 60},
    {"ingredient_id": "ING_SUATUOI", "name": "Sữa tươi thanh trùng", "base_unit": "ml", "pack_size": 1000.0, "unit_cost": 30.0, "lead_time_days": 1, "review_period_days": 3, "shelf_life_days": 10},
    {"ingredient_id": "ING_SUAGAU", "name": "Sữa gấu Thái Lan", "base_unit": "ml", "pack_size": 140.0, "unit_cost": 85.71, "lead_time_days": 3, "review_period_days": 14, "shelf_life_days": 90},
    {"ingredient_id": "ING_DUONG", "name": "Đường cát / nước đường", "base_unit": "g", "pack_size": 1000.0, "unit_cost": 25.0, "lead_time_days": 3, "review_period_days": 14, "shelf_life_days": 180},
    {"ingredient_id": "ING_KEMUOI", "name": "Kem muối béo", "base_unit": "g", "pack_size": 500.0, "unit_cost": 80.0, "lead_time_days": 2, "review_period_days": 7, "shelf_life_days": 7},
    {"ingredient_id": "ING_KEMTRUNG", "name": "Kem trứng custard", "base_unit": "g", "pack_size": 500.0, "unit_cost": 100.0, "lead_time_days": 2, "review_period_days": 7, "shelf_life_days": 7},
    {"ingredient_id": "ING_TIRAMISU", "name": "Kem phô mai Tiramisu", "base_unit": "g", "pack_size": 500.0, "unit_cost": 100.0, "lead_time_days": 2, "review_period_days": 7, "shelf_life_days": 7},
    {"ingredient_id": "ING_RICH", "name": "Kem béo Rich lùn", "base_unit": "g", "pack_size": 454.0, "unit_cost": 66.08, "lead_time_days": 2, "review_period_days": 7, "shelf_life_days": 30},
    {"ingredient_id": "ING_COTDUA", "name": "Nước cốt dừa", "base_unit": "ml", "pack_size": 400.0, "unit_cost": 67.5, "lead_time_days": 2, "review_period_days": 7, "shelf_life_days": 30},
    {"ingredient_id": "ING_MATCHA", "name": "Bột Matcha Fuji", "base_unit": "g", "pack_size": 100.0, "unit_cost": 1500.0, "lead_time_days": 3, "review_period_days": 14, "shelf_life_days": 90},
    {"ingredient_id": "ING_CACAO", "name": "Bột Cacao nguyên chất", "base_unit": "g", "pack_size": 500.0, "unit_cost": 170.0, "lead_time_days": 3, "review_period_days": 14, "shelf_life_days": 180},
    {"ingredient_id": "ING_TRA_OLONG", "name": "Cốt trà Olong", "base_unit": "ml", "pack_size": 1000.0, "unit_cost": 35.0, "lead_time_days": 2, "review_period_days": 7, "shelf_life_days": 3},
    {"ingredient_id": "ING_TRA_SUA_BASE", "name": "Cốt trà sữa truyền thống", "base_unit": "ml", "pack_size": 1000.0, "unit_cost": 20.0, "lead_time_days": 1, "review_period_days": 3, "shelf_life_days": 3},
    {"ingredient_id": "ING_MUT_TRAICAY", "name": "Mứt trái cây tổng hợp (Dâu/Ổi/Lựu)", "base_unit": "g", "pack_size": 1000.0, "unit_cost": 100.0, "lead_time_days": 3, "review_period_days": 14, "shelf_life_days": 60},
    {"ingredient_id": "ING_SIRUP_TRAICAY", "name": "Sirup trái cây (Dâu/Đào/Ổi/Vải)", "base_unit": "ml", "pack_size": 700.0, "unit_cost": 92.86, "lead_time_days": 3, "review_period_days": 14, "shelf_life_days": 90},
    {"ingredient_id": "ING_TRANCHAU_DEN", "name": "Trân châu đen nấu sẵn", "base_unit": "g", "pack_size": 1000.0, "unit_cost": 75.0, "lead_time_days": 2, "review_period_days": 7, "shelf_life_days": 1},
    {"ingredient_id": "ING_TRANCHAU_TRANG", "name": "Trân châu trắng 3Q", "base_unit": "g", "pack_size": 1000.0, "unit_cost": 60.0, "lead_time_days": 3, "review_period_days": 14, "shelf_life_days": 30},
    {"ingredient_id": "ING_COM", "name": "Cốm dẻo sữa", "base_unit": "g", "pack_size": 500.0, "unit_cost": 75.0, "lead_time_days": 2, "review_period_days": 7, "shelf_life_days": 7},
    {"ingredient_id": "ING_DABI", "name": "Đá bi sạch", "base_unit": "kg", "pack_size": 20.0, "unit_cost": 1000.0, "lead_time_days": 1, "review_period_days": 1, "shelf_life_days": 1},
    {"ingredient_id": "ING_LYNAP500", "name": "Bộ ly + nắp 500ml", "base_unit": "cái", "pack_size": 1000.0, "unit_cost": 983.0, "lead_time_days": 3, "review_period_days": 14, "shelf_life_days": 365},
    {"ingredient_id": "ING_ONGHUT", "name": "Ống hút & muỗng", "base_unit": "cái", "pack_size": 500.0, "unit_cost": 150.0, "lead_time_days": 3, "review_period_days": 14, "shelf_life_days": 365},
]


def get_real_store_item_master() -> pd.DataFrame:
    """Return DataFrame of 43 menu items with categories and base prices."""
    df = pd.DataFrame(REAL_STORE_ITEMS)
    df["launch_date"] = pd.Timestamp("2025-05-01")
    df["active"] = True
    return df


def get_real_store_ingredient_master() -> pd.DataFrame:
    """Return DataFrame of standard ingredients with pack sizes and unit costs."""
    return pd.DataFrame(REAL_STORE_INGREDIENTS)


def get_real_store_recipes() -> pd.DataFrame:
    """Return BOM recipes linking each item to ingredients with amounts and waste rates."""
    recipes = [
        # Coffee
        {"item_id": "CF_DEN", "ingredient_id": "ING_CAFE", "amount": 60.0, "unit": "ml", "waste_rate": 0.05},
        {"item_id": "CF_DEN", "ingredient_id": "ING_DUONG", "amount": 25.0, "unit": "g", "waste_rate": 0.02},
        {"item_id": "CF_DEN", "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
        {"item_id": "CF_DEN", "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        {"item_id": "CF_DEN", "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},

        {"item_id": "CF_SUA", "ingredient_id": "ING_CAFE", "amount": 60.0, "unit": "ml", "waste_rate": 0.05},
        {"item_id": "CF_SUA", "ingredient_id": "ING_SUADAC", "amount": 30.0, "unit": "g", "waste_rate": 0.03},
        {"item_id": "CF_SUA", "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
        {"item_id": "CF_SUA", "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        {"item_id": "CF_SUA", "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},

        {"item_id": "CF_BACXIU", "ingredient_id": "ING_CAFE", "amount": 30.0, "unit": "ml", "waste_rate": 0.05},
        {"item_id": "CF_BACXIU", "ingredient_id": "ING_SUADAC", "amount": 40.0, "unit": "g", "waste_rate": 0.03},
        {"item_id": "CF_BACXIU", "ingredient_id": "ING_SUATUOI", "amount": 40.0, "unit": "ml", "waste_rate": 0.03},
        {"item_id": "CF_BACXIU", "ingredient_id": "ING_RICH", "amount": 10.0, "unit": "g", "waste_rate": 0.03},
        {"item_id": "CF_BACXIU", "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
        {"item_id": "CF_BACXIU", "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        {"item_id": "CF_BACXIU", "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},

        {"item_id": "CF_MUOI", "ingredient_id": "ING_CAFE", "amount": 60.0, "unit": "ml", "waste_rate": 0.05},
        {"item_id": "CF_MUOI", "ingredient_id": "ING_SUADAC", "amount": 30.0, "unit": "g", "waste_rate": 0.03},
        {"item_id": "CF_MUOI", "ingredient_id": "ING_KEMUOI", "amount": 40.0, "unit": "g", "waste_rate": 0.05},
        {"item_id": "CF_MUOI", "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
        {"item_id": "CF_MUOI", "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        {"item_id": "CF_MUOI", "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},

        {"item_id": "CF_TRUNG", "ingredient_id": "ING_CAFE", "amount": 60.0, "unit": "ml", "waste_rate": 0.05},
        {"item_id": "CF_TRUNG", "ingredient_id": "ING_SUADAC", "amount": 20.0, "unit": "g", "waste_rate": 0.03},
        {"item_id": "CF_TRUNG", "ingredient_id": "ING_KEMTRUNG", "amount": 40.0, "unit": "g", "waste_rate": 0.05},
        {"item_id": "CF_TRUNG", "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
        {"item_id": "CF_TRUNG", "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        {"item_id": "CF_TRUNG", "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},

        {"item_id": "CF_BUONME", "ingredient_id": "ING_CAFE", "amount": 70.0, "unit": "ml", "waste_rate": 0.05},
        {"item_id": "CF_BUONME", "ingredient_id": "ING_SUADAC", "amount": 30.0, "unit": "g", "waste_rate": 0.03},
        {"item_id": "CF_BUONME", "ingredient_id": "ING_COTDUA", "amount": 30.0, "unit": "ml", "waste_rate": 0.03},
        {"item_id": "CF_BUONME", "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
        {"item_id": "CF_BUONME", "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        {"item_id": "CF_BUONME", "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},

        {"item_id": "CF_TIRAMISU", "ingredient_id": "ING_CAFE", "amount": 60.0, "unit": "ml", "waste_rate": 0.05},
        {"item_id": "CF_TIRAMISU", "ingredient_id": "ING_SUADAC", "amount": 30.0, "unit": "g", "waste_rate": 0.03},
        {"item_id": "CF_TIRAMISU", "ingredient_id": "ING_TIRAMISU", "amount": 40.0, "unit": "g", "waste_rate": 0.05},
        {"item_id": "CF_TIRAMISU", "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
        {"item_id": "CF_TIRAMISU", "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        {"item_id": "CF_TIRAMISU", "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},

        # Topping
        {"item_id": "TOP_TCDEN", "ingredient_id": "ING_TRANCHAU_DEN", "amount": 50.0, "unit": "g", "waste_rate": 0.05},
        {"item_id": "TOP_TCTRANG", "ingredient_id": "ING_TRANCHAU_TRANG", "amount": 50.0, "unit": "g", "waste_rate": 0.03},
        {"item_id": "TOP_KEMUOI", "ingredient_id": "ING_KEMUOI", "amount": 40.0, "unit": "g", "waste_rate": 0.05},
        {"item_id": "TOP_KEMTRUNG", "ingredient_id": "ING_KEMTRUNG", "amount": 40.0, "unit": "g", "waste_rate": 0.05},
        {"item_id": "TOP_KEMBUONME", "ingredient_id": "ING_TIRAMISU", "amount": 40.0, "unit": "g", "waste_rate": 0.05},

        # Milktea
        {"item_id": "MT_TRUYENTHONG", "ingredient_id": "ING_TRA_SUA_BASE", "amount": 250.0, "unit": "ml", "waste_rate": 0.03},
        {"item_id": "MT_TRUYENTHONG", "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
        {"item_id": "MT_TRUYENTHONG", "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        {"item_id": "MT_TRUYENTHONG", "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},

        {"item_id": "MT_DUONGDEN", "ingredient_id": "ING_SUATUOI", "amount": 200.0, "unit": "ml", "waste_rate": 0.03},
        {"item_id": "MT_DUONGDEN", "ingredient_id": "ING_TRANCHAU_DEN", "amount": 60.0, "unit": "g", "waste_rate": 0.05},
        {"item_id": "MT_DUONGDEN", "ingredient_id": "ING_DUONG", "amount": 20.0, "unit": "g", "waste_rate": 0.02},
        {"item_id": "MT_DUONGDEN", "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
        {"item_id": "MT_DUONGDEN", "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        {"item_id": "MT_DUONGDEN", "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},

        {"item_id": "MT_THAI", "ingredient_id": "ING_TRA_SUA_BASE", "amount": 250.0, "unit": "ml", "waste_rate": 0.03},
        {"item_id": "MT_THAI", "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
        {"item_id": "MT_THAI", "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        {"item_id": "MT_THAI", "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},

        {"item_id": "MT_DAU_VQ", "ingredient_id": "ING_TRA_SUA_BASE", "amount": 200.0, "unit": "ml", "waste_rate": 0.03},
        {"item_id": "MT_DAU_VQ", "ingredient_id": "ING_SIRUP_TRAICAY", "amount": 25.0, "unit": "ml", "waste_rate": 0.03},
        {"item_id": "MT_DAU_VQ", "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
        {"item_id": "MT_DAU_VQ", "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        {"item_id": "MT_DAU_VQ", "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},

        {"item_id": "MT_MATCHA", "ingredient_id": "ING_TRA_SUA_BASE", "amount": 200.0, "unit": "ml", "waste_rate": 0.03},
        {"item_id": "MT_MATCHA", "ingredient_id": "ING_MATCHA", "amount": 3.0, "unit": "g", "waste_rate": 0.05},
        {"item_id": "MT_MATCHA", "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
        {"item_id": "MT_MATCHA", "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        {"item_id": "MT_MATCHA", "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},

        {"item_id": "MT_KHOAIMON", "ingredient_id": "ING_TRA_SUA_BASE", "amount": 220.0, "unit": "ml", "waste_rate": 0.03},
        {"item_id": "MT_KHOAIMON", "ingredient_id": "ING_SIRUP_TRAICAY", "amount": 20.0, "unit": "ml", "waste_rate": 0.03},
        {"item_id": "MT_KHOAIMON", "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
        {"item_id": "MT_KHOAIMON", "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        {"item_id": "MT_KHOAIMON", "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},

        {"item_id": "MT_TAOXANH", "ingredient_id": "ING_TRA_SUA_BASE", "amount": 220.0, "unit": "ml", "waste_rate": 0.03},
        {"item_id": "MT_TAOXANH", "ingredient_id": "ING_SIRUP_TRAICAY", "amount": 20.0, "unit": "ml", "waste_rate": 0.03},
        {"item_id": "MT_TAOXANH", "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
        {"item_id": "MT_TAOXANH", "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        {"item_id": "MT_TAOXANH", "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},

        {"item_id": "MT_COMDEO", "ingredient_id": "ING_TRA_SUA_BASE", "amount": 200.0, "unit": "ml", "waste_rate": 0.03},
        {"item_id": "MT_COMDEO", "ingredient_id": "ING_COM", "amount": 30.0, "unit": "g", "waste_rate": 0.05},
        {"item_id": "MT_COMDEO", "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
        {"item_id": "MT_COMDEO", "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        {"item_id": "MT_COMDEO", "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},

        {"item_id": "MT_OLONG_LAI", "ingredient_id": "ING_TRA_OLONG", "amount": 150.0, "unit": "ml", "waste_rate": 0.03},
        {"item_id": "MT_OLONG_LAI", "ingredient_id": "ING_SUATUOI", "amount": 80.0, "unit": "ml", "waste_rate": 0.03},
        {"item_id": "MT_OLONG_LAI", "ingredient_id": "ING_SUADAC", "amount": 25.0, "unit": "g", "waste_rate": 0.03},
        {"item_id": "MT_OLONG_LAI", "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
        {"item_id": "MT_OLONG_LAI", "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        {"item_id": "MT_OLONG_LAI", "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},

        # Tea_Soda (11 items)
        {"item_id": "TEA_HONGTRA", "ingredient_id": "ING_TRA_OLONG", "amount": 150.0, "unit": "ml", "waste_rate": 0.03},
        {"item_id": "TEA_HONGTRA", "ingredient_id": "ING_DUONG", "amount": 25.0, "unit": "g", "waste_rate": 0.02},
        {"item_id": "TEA_HONGTRA", "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
        {"item_id": "TEA_HONGTRA", "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        {"item_id": "TEA_HONGTRA", "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
    ]

    # Dynamically fill general tea/soda, iceblend, matcha/cacao items
    fruit_teas = [
        "TEA_DAU", "TEA_VIETQUAT", "TEA_KIWI", "TEA_CHANHDAY", "TEA_LUUDO",
        "TEA_OIHONG", "TEA_TAOXANH", "TEA_DAOCAMSA", "TEA_DUALUOI", "TEA_MANGCAU"
    ]
    for tid in fruit_teas:
        recipes.extend([
            {"item_id": tid, "ingredient_id": "ING_TRA_OLONG", "amount": 120.0, "unit": "ml", "waste_rate": 0.03},
            {"item_id": tid, "ingredient_id": "ING_MUT_TRAICAY", "amount": 30.0, "unit": "g", "waste_rate": 0.05},
            {"item_id": tid, "ingredient_id": "ING_SIRUP_TRAICAY", "amount": 15.0, "unit": "ml", "waste_rate": 0.03},
            {"item_id": tid, "ingredient_id": "ING_DABI", "amount": 0.2, "unit": "kg", "waste_rate": 0.05},
            {"item_id": tid, "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
            {"item_id": tid, "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        ])

    iceblends = ["ICE_CF_CARAMEL", "ICE_TRAICAY", "ICE_MATCHA", "ICE_OREO", "ICE_MILO"]
    for iid in iceblends:
        recipes.extend([
            {"item_id": iid, "ingredient_id": "ING_SUATUOI", "amount": 100.0, "unit": "ml", "waste_rate": 0.03},
            {"item_id": iid, "ingredient_id": "ING_SUADAC", "amount": 30.0, "unit": "g", "waste_rate": 0.03},
            {"item_id": iid, "ingredient_id": "ING_DABI", "amount": 0.3, "unit": "kg", "waste_rate": 0.08},
            {"item_id": iid, "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
            {"item_id": iid, "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01},
        ])

    matcha_cacao = ["MAT_LATTE", "MAT_GAU", "MAT_FRUIT", "CAC_LATTE", "CAC_GAU", "DES_PANNACOTTA"]
    for mid in matcha_cacao:
        if "MAT" in mid:
            recipes.append({"item_id": mid, "ingredient_id": "ING_MATCHA", "amount": 4.0, "unit": "g", "waste_rate": 0.05})
            recipes.append({"item_id": mid, "ingredient_id": "ING_SUATUOI", "amount": 100.0, "unit": "ml", "waste_rate": 0.03})
            recipes.append({"item_id": mid, "ingredient_id": "ING_SUADAC", "amount": 25.0, "unit": "g", "waste_rate": 0.03})
        elif "CAC" in mid:
            recipes.append({"item_id": mid, "ingredient_id": "ING_CACAO", "amount": 8.0, "unit": "g", "waste_rate": 0.05})
            recipes.append({"item_id": mid, "ingredient_id": "ING_SUATUOI", "amount": 100.0, "unit": "ml", "waste_rate": 0.03})
            recipes.append({"item_id": mid, "ingredient_id": "ING_SUADAC", "amount": 25.0, "unit": "g", "waste_rate": 0.03})
        else: # Panna cotta
            recipes.append({"item_id": mid, "ingredient_id": "ING_SUATUOI", "amount": 60.0, "unit": "ml", "waste_rate": 0.02})
            recipes.append({"item_id": mid, "ingredient_id": "ING_SUADAC", "amount": 20.0, "unit": "g", "waste_rate": 0.02})
            recipes.append({"item_id": mid, "ingredient_id": "ING_RICH", "amount": 20.0, "unit": "g", "waste_rate": 0.02})
        
        recipes.append({"item_id": mid, "ingredient_id": "ING_LYNAP500", "amount": 1.0, "unit": "cái", "waste_rate": 0.01})
        recipes.append({"item_id": mid, "ingredient_id": "ING_ONGHUT", "amount": 1.0, "unit": "cái", "waste_rate": 0.01})

    return pd.DataFrame(recipes)


def parse_real_store_excel(excel_path: Path) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """
    Parse Tháng 5 and Tháng 6 from dataset.xlsx to extract daily revenue and ice costs.
    Handles both raw input format and calibrated multi-tab format.
    """
    wb = openpyxl.load_workbook(str(excel_path), data_only=True)
    
    rev_records = []
    ice_records = []
    
    if "Tháng 5" in wb.sheetnames:
        # Raw dataset format
        s5 = wb["Tháng 5"]
        for day in range(1, 32):
            col = day + 1  # Column 2 is Day 1
            rev = s5.cell(row=8, column=col).value
            ice = s5.cell(row=13, column=col).value
            
            if rev is not None and float(rev) > 0:
                date_str = f"2025-05-{day:02d}"
                rev_records.append({"date": pd.Timestamp(date_str), "revenue": float(rev)})
            if ice is not None and float(ice) > 0:
                date_str = f"2025-05-{day:02d}"
                ice_records.append({"date": pd.Timestamp(date_str), "ice_cost": float(ice)})
                
        s6 = wb["Tháng 6"]
        for day in range(1, 31):
            col = day + 1
            rev = s6.cell(row=8, column=col).value
            ice = s6.cell(row=13, column=col).value
            
            if rev is not None and float(rev) > 0:
                date_str = f"2025-06-{day:02d}"
                rev_records.append({"date": pd.Timestamp(date_str), "revenue": float(rev)})
                
            if ice is not None and float(ice) > 0:
                date_str = f"2025-06-{day:02d}"
                ice_records.append({"date": pd.Timestamp(date_str), "ice_cost": float(ice)})
    else:
        # Multi-tab monthly format (e.g. dataset_14_thang_thuc_te.xlsx)
        s5 = wb["Tháng 5_2025"]
        for day in range(1, 32):
            col = day + 4  # Col 5 (E) is Day 1
            rev = s5.cell(row=49, column=col).value
            if rev is not None and float(rev) > 0:
                date_str = f"2025-05-{day:02d}"
                rev_records.append({"date": pd.Timestamp(date_str), "revenue": float(rev)})
                ice_records.append({"date": pd.Timestamp(date_str), "ice_cost": 20000.0})
                
        s6 = wb["Tháng 6_2025"]
        for day in range(1, 31):
            col = day + 4
            rev = s6.cell(row=49, column=col).value
            if rev is not None and float(rev) > 0:
                date_str = f"2025-06-{day:02d}"
                rev_records.append({"date": pd.Timestamp(date_str), "revenue": float(rev)})
                ice_records.append({"date": pd.Timestamp(date_str), "ice_cost": 20000.0})
            
    df_rev = pd.DataFrame(rev_records).sort_values("date").reset_index(drop=True)
    df_ice = pd.DataFrame(ice_records).sort_values("date").reset_index(drop=True)
    
    return df_rev, df_ice
